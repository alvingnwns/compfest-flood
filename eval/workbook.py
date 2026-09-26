from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_CANDIDATES = (
    REPO_ROOT / "data" / "ARUNA_Dummy_Company_Test_Data.xlsx",
    REPO_ROOT / "ARUNA_Dummy_Company_Test_Data.xlsx",
)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SHEET_COLUMNS: dict[str, tuple[str, ...]] = {
    "Products": ("productId", "productName", "sellingPrice", "unit"),
    "Orders": ("orderId", "storeId", "productId", "quantity", "priority", "deadlineMinutes"),
    "Inventory": ("warehouseId", "productId", "availableQuantity"),
    "Materials": ("materialId", "materialName", "supplierId", "availableQuantity"),
    "BOM": ("productId", "materialId", "quantityRequired"),
}

BusinessData = dict[str, list[dict[str, Any]]]

# Copy of ARUNA_Dummy_Company_Test_Data.xlsx, used only when the file is absent.
_FALLBACK_ROWS: dict[str, list[tuple[Any, ...]]] = {
    "Products": [
        ("P001", "Frozen Chicken 1kg", 62000, "pack"),
        ("P002", "Chicken Nugget 500g", 48000, "pack"),
        ("P003", "Beef Meatball 500g", 55000, "pack"),
        ("P004", "Fish Fillet 500g", 72000, "pack"),
        ("P005", "Shrimp Tempura 400g", 68000, "pack"),
        ("P006", "French Fries 1kg", 45000, "pack"),
        ("P007", "Mixed Vegetables 500g", 32000, "pack"),
        ("P008", "Chicken Sausage 500g", 52000, "pack"),
        ("P009", "Beef Sausage 500g", 59000, "pack"),
        ("P010", "Chicken Karaage 500g", 56000, "pack"),
        ("P011", "Fish Ball 500g", 43000, "pack"),
        ("P012", "Seafood Dumpling 500g", 51000, "pack"),
        ("P013", "Spicy Chicken Wings 700g", 65000, "pack"),
        ("P014", "Chicken Patty 600g", 58000, "pack"),
        ("P015", "Ready Meal Teriyaki 350g", 39000, "box"),
    ],
    "Orders": [
        ("ORDER-001", "store-a", "P001", 180, "critical", 25),
        ("ORDER-002", "store-b", "P002", 140, "high", 30),
        ("ORDER-003", "store-c", "P003", 120, "high", 35),
        ("ORDER-004", "store-d", "P004", 100, "critical", 25),
        ("ORDER-005", "store-e", "P005", 80, "normal", 50),
        ("ORDER-006", "store-a", "P006", 160, "normal", 45),
        ("ORDER-007", "store-b", "P007", 130, "normal", 55),
        ("ORDER-008", "store-c", "P008", 110, "high", 35),
        ("ORDER-009", "store-d", "P009", 90, "critical", 30),
        ("ORDER-010", "store-e", "P010", 120, "high", 40),
        ("ORDER-011", "store-a", "P011", 140, "normal", 50),
        ("ORDER-012", "store-b", "P012", 100, "high", 40),
        ("ORDER-013", "store-c", "P013", 95, "critical", 25),
        ("ORDER-014", "store-d", "P014", 105, "normal", 60),
        ("ORDER-015", "store-e", "P015", 150, "high", 45),
        ("ORDER-016", "store-a", "P002", 90, "critical", 25),
        ("ORDER-017", "store-b", "P004", 75, "high", 35),
        ("ORDER-018", "store-c", "P006", 125, "normal", 55),
        ("ORDER-019", "store-d", "P010", 85, "high", 40),
        ("ORDER-020", "store-e", "P001", 130, "critical", 30),
    ],
    "Inventory": [
        ("wh-west", "P001", 75), ("wh-east", "P001", 55),
        ("wh-west", "P002", 80), ("wh-east", "P002", 60),
        ("wh-west", "P003", 45), ("wh-east", "P003", 40),
        ("wh-west", "P004", 35), ("wh-east", "P004", 45),
        ("wh-west", "P005", 50), ("wh-east", "P005", 35),
        ("wh-west", "P006", 100), ("wh-east", "P006", 70),
        ("wh-west", "P007", 85), ("wh-east", "P007", 75),
        ("wh-west", "P008", 55), ("wh-east", "P008", 50),
        ("wh-west", "P009", 40), ("wh-east", "P009", 35),
        ("wh-west", "P010", 60), ("wh-east", "P010", 50),
        ("wh-west", "P011", 70), ("wh-east", "P011", 55),
        ("wh-west", "P012", 45), ("wh-east", "P012", 45),
        ("wh-west", "P013", 35), ("wh-east", "P013", 30),
        ("wh-west", "P014", 50), ("wh-east", "P014", 40),
        ("wh-west", "P015", 85), ("wh-east", "P015", 75),
    ],
    "Materials": [
        ("M001", "Raw Chicken", "sup-a", 1150),
        ("M002", "Raw Beef", "sup-a", 420),
        ("M003", "Raw Fish", "sup-a", 520),
        ("M004", "Raw Shrimp", "sup-a", 280),
        ("M005", "Frozen Potato", "sup-a", 720),
        ("M006", "Mixed Vegetables Raw", "sup-a", 500),
        ("M007", "Seasoning Mix", "sup-b", 380),
        ("M008", "Bread Crumbs", "sup-b", 260),
        ("M009", "Flour", "sup-b", 340),
        ("M010", "Cooking Oil", "sup-b", 300),
        ("M011", "Sausage Casing", "sup-b", 190),
        ("M012", "Batter Mix", "sup-b", 250),
        ("M013", "Sauce Base", "sup-b", 280),
        ("M014", "Plastic Pouch", "sup-b", 1450),
        ("M015", "Carton Box", "sup-b", 800),
        ("M016", "Label Sticker", "sup-b", 1800),
        ("M017", "Starch Binder", "sup-b", 240),
        ("M018", "Rice", "sup-a", 520),
    ],
    "BOM": [
        ("P001", "M001", 1), ("P001", "M007", 0.05), ("P001", "M014", 0.02), ("P001", "M016", 0.02),
        ("P002", "M001", 0.6), ("P002", "M008", 0.1), ("P002", "M009", 0.05), ("P002", "M007", 0.04),
        ("P002", "M014", 0.02),
        ("P003", "M002", 0.6), ("P003", "M017", 0.08), ("P003", "M007", 0.04), ("P003", "M014", 0.02),
        ("P004", "M003", 0.55), ("P004", "M007", 0.03), ("P004", "M014", 0.02), ("P004", "M016", 0.02),
        ("P005", "M004", 0.45), ("P005", "M012", 0.1), ("P005", "M010", 0.02), ("P005", "M014", 0.02),
        ("P006", "M005", 1), ("P006", "M007", 0.02), ("P006", "M014", 0.02),
        ("P007", "M006", 0.55), ("P007", "M014", 0.02), ("P007", "M016", 0.02),
        ("P008", "M001", 0.5), ("P008", "M011", 0.03), ("P008", "M007", 0.03), ("P008", "M014", 0.02),
        ("P009", "M002", 0.5), ("P009", "M011", 0.03), ("P009", "M007", 0.03), ("P009", "M014", 0.02),
        ("P010", "M001", 0.55), ("P010", "M012", 0.08), ("P010", "M007", 0.03), ("P010", "M014", 0.02),
        ("P011", "M003", 0.4), ("P011", "M017", 0.06), ("P011", "M007", 0.02), ("P011", "M014", 0.02),
        ("P012", "M003", 0.3), ("P012", "M004", 0.15), ("P012", "M009", 0.05), ("P012", "M007", 0.03),
        ("P012", "M014", 0.02),
        ("P013", "M001", 0.7), ("P013", "M013", 0.08), ("P013", "M007", 0.03), ("P013", "M014", 0.02),
        ("P014", "M001", 0.6), ("P014", "M017", 0.08), ("P014", "M007", 0.03), ("P014", "M014", 0.02),
        ("P015", "M001", 0.2), ("P015", "M018", 0.15), ("P015", "M006", 0.1), ("P015", "M013", 0.06),
        ("P015", "M015", 0.03), ("P015", "M016", 0.02),
    ],
}


def fallback_data() -> BusinessData:
    return {
        sheet: [dict(zip(SHEET_COLUMNS[sheet], row, strict=True)) for row in rows]
        for sheet, rows in _FALLBACK_ROWS.items()
    }


def load_nominal() -> tuple[BusinessData, str]:
    for path in SOURCE_CANDIDATES:
        if path.exists():
            return _read_workbook(path), path.relative_to(REPO_ROOT).as_posix()
    return fallback_data(), "hardcoded-fallback"


def _read_workbook(path: Path) -> BusinessData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    data: BusinessData = {}
    for sheet, columns in SHEET_COLUMNS.items():
        rows = workbook[sheet].iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        data[sheet] = [
            {column: values[headers.index(column)] if column in headers else None for column in columns}
            for values in rows
            if any(value is not None and str(value).strip() for value in values)
        ]
    workbook.close()
    return data


def build_workbook(data: BusinessData) -> bytes:
    """Serialize business data to .xlsx; sheets absent from `data` are omitted from the file."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet, columns in SHEET_COLUMNS.items():
        if sheet not in data:
            continue
        worksheet = workbook.create_sheet(sheet)
        worksheet.append(list(columns))
        for row in data[sheet]:
            worksheet.append([row.get(column) for column in columns])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
